import sys
import pandas as pd
import re

from PyQt6.QtWidgets import (
    QApplication, QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QLabel,
    QTableWidget, QTableWidgetItem, QMessageBox, QFileDialog,
    QTabWidget, QComboBox, QLineEdit, QListWidget
)
from PyQt6.QtCore import Qt
from datetime import datetime, timedelta, time
from PyQt6.QtGui import QColor
from openpyxl.styles import PatternFill
from aplicacion_presentismo import AplicacionPresentismo


# ============================================
# 🟢 HELPERS GENERALES
# ============================================

def _norm_col(s: str) -> str:
    s = str(s).strip().lower()
    s = s.replace("\n", " ").replace("\r", " ")
    s = re.sub(r"\s+", " ", s)
    return s


def _find_col(df, candidates_exact=(), candidates_contains=()):
    cols = list(df.columns)
    norm_map = {c: _norm_col(c) for c in cols}

    exact_norm = {_norm_col(x) for x in candidates_exact}
    for c, nc in norm_map.items():
        if nc in exact_norm:
            return c

    contains_norm = [_norm_col(x) for x in candidates_contains]
    for c, nc in norm_map.items():
        for pat in contains_norm:
            if pat and pat in nc:
                return c

    return None


def cargar_excel_robusto(ruta: str) -> pd.DataFrame:
    """Lee xlsx/xls/csv de forma robusta y devuelve un DataFrame."""
    ruta_lower = str(ruta).lower().strip()

    if ruta_lower.endswith(".csv"):
        try:
            return pd.read_csv(ruta, encoding="utf-8")
        except Exception:
            return pd.read_csv(ruta, encoding="latin-1")

    try:
        return pd.read_excel(ruta, engine="openpyxl")
    except Exception:
        pass

    try:
        return pd.read_excel(ruta, engine="xlrd")
    except Exception:
        pass

    return pd.read_excel(ruta)


def detectar_columnas_reloj(df):
    col_nombre = _find_col(
        df,
        candidates_exact=("Nombre", "Empleado", "Personal", "Funcionario"),
        candidates_contains=("nombre", "emple", "funcion", "personal"),
    )

    col_fechahora = _find_col(
        df,
        candidates_exact=("Fecha/Hora", "Fecha Hora", "Fecha y Hora", "Marc.", "Marcación", "Marcacion", "Timestamp"),
        candidates_contains=("fecha/hora", "fecha hora", "marc", "timestamp", "fecha y hora"),
    )

    col_tipo = _find_col(
        df,
        candidates_exact=("Tipo de registro", "Tipo", "Evento", "Movimiento"),
        candidates_contains=("tipo de registro", "tipo", "evento", "movim"),
    )

    return col_nombre, col_fechahora, col_tipo


def _to_timedelta_series(s: pd.Series) -> pd.Series:
    td = pd.to_timedelta(s.astype(str), errors="coerce")
    return td.fillna(pd.Timedelta(0))


# ============================================
# 🟡 PESTAÑA BASE SIMPLE
# ============================================
class PestañaExcel(QWidget):
    """Clase base para pestañas simples de carga de Excel"""

    def __init__(self, titulo):
        super().__init__()
        self.titulo = titulo
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()
        self.label = QLabel(f"📂 {self.titulo}")
        self.label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self.label)

        self.boton_cargar = QPushButton("Seleccionar archivo Excel")
        self.boton_cargar.clicked.connect(self.cargar_excel)
        layout.addWidget(self.boton_cargar)

        self.tabla = QTableWidget()
        layout.addWidget(self.tabla)

        self.setLayout(layout)

    def cargar_excel(self):
        archivo, _ = QFileDialog.getOpenFileName(
            self, "Seleccionar archivo Excel", "", "Archivos Excel (*.xlsx *.xls *.csv)"
        )
        if archivo:
            try:
                df = cargar_excel_robusto(archivo)
                self.mostrar_en_tabla(df)
            except Exception as e:
                QMessageBox.critical(self, "Error", f"No se pudo leer el archivo:\n{e}")

    def mostrar_en_tabla(self, df):
        self.tabla.setRowCount(0)
        self.tabla.setColumnCount(0)
        if df.empty:
            QMessageBox.warning(self, "Atención", "El archivo Excel está vacío.")
            return
        self.tabla.setColumnCount(len(df.columns))
        self.tabla.setHorizontalHeaderLabels(df.columns.astype(str).tolist())
        for i, fila in df.iterrows():
            self.tabla.insertRow(i)
            for j, valor in enumerate(fila):
                item = QTableWidgetItem(str(valor))
                self.tabla.setItem(i, j, item)
        self.label.setText(f"✅ {self.titulo} - Archivo cargado correctamente ({len(df)} filas)")


# ============================================
# 🔵 APLICACIÓN 1 - CONTROL DE FICHAJE Y HORAS EXTRAS
# ============================================
class Aplicacion1(QWidget):
    """Procesa fichajes: limpia, agrupa, calcula horas trabajadas y horas extras."""

    def __init__(self):
        super().__init__()
        self.df_procesado = None
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()
        self.label = QLabel("📊 Aplicación 1: Control de fichaje y horas extras")
        self.label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self.label)

        self.boton_cargar = QPushButton("Seleccionar archivo Excel de fichajes")
        self.boton_cargar.clicked.connect(self.cargar_y_procesar_excel)
        layout.addWidget(self.boton_cargar)

        self.boton_exportar = QPushButton("Exportar resultados a Excel")
        self.boton_exportar.clicked.connect(self.exportar_excel)
        layout.addWidget(self.boton_exportar)

        filtro_layout = QHBoxLayout()
        self.combo_empleados = QComboBox()
        self.combo_empleados.currentIndexChanged.connect(self.filtrar_empleado)
        self.boton_mostrar_todos = QPushButton("Mostrar todos")
        self.boton_mostrar_todos.clicked.connect(self.mostrar_todos)
        filtro_layout.addWidget(QLabel("Filtrar por empleado:"))
        filtro_layout.addWidget(self.combo_empleados)
        filtro_layout.addWidget(self.boton_mostrar_todos)
        layout.addLayout(filtro_layout)

        eliminar_layout = QHBoxLayout()
        self.lista_dias = QListWidget()
        self.lista_dias.setSelectionMode(QListWidget.SelectionMode.MultiSelection)
        for d in ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]:
            self.lista_dias.addItem(d)
        eliminar_layout.addWidget(QLabel("Eliminar días de la semana:"))
        eliminar_layout.addWidget(self.lista_dias)

        self.input_fechas = QLineEdit()
        self.input_fechas.setPlaceholderText("Fechas a eliminar: dd/mm/yyyy, dd/mm/yyyy...")
        eliminar_layout.addWidget(self.input_fechas)

        self.boton_eliminar = QPushButton("Aplicar eliminación")
        self.boton_eliminar.clicked.connect(self.eliminar_dias_fechas)
        eliminar_layout.addWidget(self.boton_eliminar)
        layout.addLayout(eliminar_layout)

        self.tabla = QTableWidget()
        layout.addWidget(self.tabla)
        self.setLayout(layout)

    def cargar_y_procesar_excel(self):
        archivo, _ = QFileDialog.getOpenFileName(
            self,
            "Seleccionar archivo",
            "",
            "Archivos (*.xlsx *.xls *.csv)",
        )
        if not archivo:
            return

        try:
            df_raw = cargar_excel_robusto(archivo)
            self._raw_rows = len(df_raw)

            self.df_procesado = self.procesar_fichajes(df_raw)

            self.combo_empleados.clear()
            empleados = sorted(self.df_procesado["Nombre"].dropna().unique())
            self.combo_empleados.addItems(["-- Seleccione --"] + list(empleados))

            self.mostrar_en_tabla(self.df_procesado)

            raw = getattr(self, "_raw_rows", 0)
            valid = getattr(self, "_valid_rows", 0)
            completos = getattr(self, "_dias_completos", 0)
            incompletos = getattr(self, "_dias_incompletos", 0)
            sinfichaje = getattr(self, "_dias_sin", 0)

            self.label.setText(
                f"✅ Procesado | Filas crudas: {raw} | Filas válidas: {valid} | "
                f"Días completos: {completos} | Incompletos: {incompletos} | Sin fichaje: {sinfichaje}"
            )

        except Exception as e:
            QMessageBox.critical(
                self,
                "Error",
                f"No se pudo procesar el archivo:\n{e}"
            )

    def procesar_fichajes(self, df):
        col_nombre, col_fechahora, col_tipo = detectar_columnas_reloj(df)

        if not col_nombre:
            raise Exception("No pude detectar la columna de NOMBRE (ej: 'Nombre' / 'Empleado').")
        if not col_fechahora:
            raise Exception("No pude detectar la columna de FECHA/HORA (ej: 'Fecha/Hora' / 'Marc.' / 'Timestamp').")

        df = df.copy()
        df["Nombre"] = df[col_nombre].astype(str).str.strip()
        df["Fecha/Hora"] = pd.to_datetime(df[col_fechahora], errors="coerce")
        df = df.dropna(subset=["Fecha/Hora"])
        self._valid_rows = len(df)

        df["Fecha"] = df["Fecha/Hora"].dt.date
        df["Hora"] = df["Fecha/Hora"].dt.time

        if col_tipo and col_tipo in df.columns:
            df["Tipo de registro"] = df[col_tipo].astype(str).str.strip().str.lower()

        resultado = []
        for (nombre, fecha), grupo in df.groupby(["Nombre", "Fecha"]):
            grupo_ordenado = grupo.sort_values("Fecha/Hora")
            hora_entrada = grupo_ordenado.iloc[0]["Hora"] if not grupo_ordenado.empty else None
            hora_salida = grupo_ordenado.iloc[-1]["Hora"] if len(grupo_ordenado) > 1 else None

            if len(grupo_ordenado) == 0:
                entrada = salida = total_horas = horas_extra = ""
                estado = "Sin fichaje"
            elif len(grupo_ordenado) == 1:
                if grupo_ordenado.iloc[0]["Hora"].hour < 12:
                    entrada = hora_entrada.strftime("%H:%M:%S")
                    salida = total_horas = horas_extra = ""
                    estado = "⚠️ Falta salida"
                else:
                    entrada = total_horas = horas_extra = ""
                    salida = hora_entrada.strftime("%H:%M:%S")
                    estado = "⚠️ Falta entrada"
            else:
                entrada = hora_entrada.strftime("%H:%M:%S") if hora_entrada else ""
                salida = hora_salida.strftime("%H:%M:%S") if hora_salida else ""
                estado = "✅ Completo"

                tiempo_entrada = datetime.combine(datetime.today(), hora_entrada)
                tiempo_salida = datetime.combine(datetime.today(), hora_salida)
                total = tiempo_salida - tiempo_entrada
                limite = timedelta(hours=8)
                extra = total - limite if total > limite else timedelta(0)

                total_horas = str(total).split(".")[0]
                horas_extra = str(extra).split(".")[0]

            resultado.append(
                {
                    "Nombre": nombre,
                    "Fecha": fecha.strftime("%d/%m/%Y"),
                    "Entrada": entrada,
                    "Salida": salida,
                    "Estado": estado,
                    "Total Horas": total_horas,
                    "Horas Extra": horas_extra,
                }
            )

        self._dias_completos = sum(1 for r in resultado if r.get("Estado") == "✅ Completo")
        self._dias_incompletos = sum(1 for r in resultado if r.get("Estado") and "⚠️" in str(r.get("Estado")))
        self._dias_sin = sum(1 for r in resultado if r.get("Estado") == "Sin fichaje")

        return pd.DataFrame(resultado)

    def mostrar_en_tabla(self, df):
        self.tabla.setRowCount(0)
        self.tabla.setColumnCount(0)
        if df is None or df.empty:
            QMessageBox.warning(self, "Atención", "No se encontraron fichajes válidos.")
            return

        self.tabla.setColumnCount(len(df.columns))
        self.tabla.setHorizontalHeaderLabels(df.columns.astype(str).tolist())

        for i, fila in df.iterrows():
            self.tabla.insertRow(i)
            for j, valor in enumerate(fila):
                self.tabla.setItem(i, j, QTableWidgetItem(str(valor)))

            color = None
            estado = str(fila.get("Estado", ""))
            if estado == "✅ Completo":
                color = QColor(200, 255, 200)
            elif "⚠️" in estado:
                color = QColor(255, 255, 150)
            elif estado == "Sin fichaje":
                color = QColor(255, 200, 200)

            if color:
                for j in range(len(fila)):
                    item = self.tabla.item(i, j)
                    if item:
                        item.setBackground(color)

    def filtrar_empleado(self):
        if self.df_procesado is None or self.combo_empleados.currentIndex() <= 0:
            return

        nombre = self.combo_empleados.currentText()
        df_filtrado = self.df_procesado[self.df_procesado["Nombre"] == nombre].copy()

        ser = _to_timedelta_series(df_filtrado["Horas Extra"])
        total_extra = str(timedelta(seconds=int(ser.dt.total_seconds().sum())))

        if not df_filtrado.empty:
            df_filtrado.loc[len(df_filtrado)] = {
                "Nombre": f"TOTAL DE HORAS EXTRAS DE {nombre}",
                "Fecha": "",
                "Entrada": "",
                "Salida": "",
                "Estado": "",
                "Total Horas": "",
                "Horas Extra": total_extra,
            }

        self.mostrar_en_tabla(df_filtrado)

    def mostrar_todos(self):
        if self.df_procesado is not None:
            self.mostrar_en_tabla(self.df_procesado)

    def eliminar_dias_fechas(self):
        if self.df_procesado is None:
            return

        df = self.df_procesado.copy()

        dias_seleccionados = [item.text() for item in self.lista_dias.selectedItems()]
        if dias_seleccionados:
            dias_map = {
                "Lunes": 0,
                "Martes": 1,
                "Miércoles": 2,
                "Jueves": 3,
                "Viernes": 4,
                "Sábado": 5,
                "Domingo": 6,
            }
            wanted = {dias_map[d] for d in dias_seleccionados}
            idx_drop = []
            for i, row in df.iterrows():
                try:
                    fecha_obj = datetime.strptime(str(row["Fecha"]), "%d/%m/%Y").date()
                except Exception:
                    continue
                if fecha_obj.weekday() in wanted:
                    idx_drop.append(i)
            df = df.drop(idx_drop)

        fechas_input = self.input_fechas.text().strip()
        if fechas_input:
            fechas = []
            for f in fechas_input.split(","):
                f = f.strip()
                if not f:
                    continue
                try:
                    fechas.append(datetime.strptime(f, "%d/%m/%Y").date())
                except Exception:
                    pass
            fechas_set = set(fechas)
            df = df[
                ~df["Fecha"].apply(
                    lambda x: datetime.strptime(str(x), "%d/%m/%Y").date() in fechas_set if str(x).strip() else False
                )
            ]

        self.df_procesado = df.reset_index(drop=True)
        self.mostrar_en_tabla(self.df_procesado)

    def exportar_excel(self):
        if self.df_procesado is None or self.df_procesado.empty:
            QMessageBox.warning(self, "Atención", "No hay datos para exportar.")
            return

        archivo, _ = QFileDialog.getSaveFileName(
            self, "Guardar Excel", "", "Archivos Excel (*.xlsx)"
        )
        if not archivo:
            return

        try:
            with pd.ExcelWriter(archivo, engine="openpyxl") as writer:
                df_export = self.df_procesado.copy()

                df_final = pd.DataFrame()
                for nombre, grupo in df_export.groupby("Nombre"):
                    df_final = pd.concat([df_final, grupo], ignore_index=True)

                    td_extra = _to_timedelta_series(grupo["Horas Extra"]).sum()
                    if td_extra > pd.Timedelta(0):
                        df_final = pd.concat(
                            [
                                df_final,
                                pd.DataFrame(
                                    [
                                        {
                                            "Nombre": f"TOTAL HORAS EXTRAS DE {nombre}",
                                            "Fecha": "",
                                            "Entrada": "",
                                            "Salida": "",
                                            "Estado": "",
                                            "Total Horas": "",
                                            "Horas Extra": str(td_extra),
                                        }
                                    ]
                                ),
                            ],
                            ignore_index=True,
                        )

                df_final.to_excel(writer, index=False, sheet_name="Datos Procesados")

                ws = writer.sheets["Datos Procesados"]
                amarillo = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
                for i, row in enumerate(df_final.itertuples(), start=2):
                    if "TOTAL HORAS EXTRAS" in str(row.Nombre):
                        for j in range(1, len(df_final.columns) + 1):
                            ws.cell(row=i, column=j).fill = amarillo

                tmp = df_export.copy()
                tmp["_td_extra"] = _to_timedelta_series(tmp["Horas Extra"])
                resumen = (
                    tmp.groupby("Nombre", as_index=False)["_td_extra"]
                    .sum()
                    .rename(columns={"_td_extra": "Total Horas Extra"})
                )
                resumen["Total Horas Extra"] = resumen["Total Horas Extra"].astype(str)
                resumen.to_excel(writer, index=False, sheet_name="Total Horas Extra")

            QMessageBox.information(self, "Éxito", "Archivo exportado correctamente.")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo exportar el archivo:\n{e}")


# ============================================
# 🟢 APLICACIÓN 2 - HORAS EXTRA ANTES DE 06 AM
# ============================================
class Aplicacion2(QWidget):
    """Procesa la primera marca del día por persona y calcula horas antes de las 06:00 AM."""
    def __init__(self):
        super().__init__()
        self.df_procesado = None
        self.df_base = None
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()

        self.label = QLabel("📊 Aplicación 2: Horas antes de las 06 AM")
        self.label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self.label)

        self.boton_cargar = QPushButton("Seleccionar Excel")
        self.boton_cargar.clicked.connect(self.cargar_y_procesar_excel)
        layout.addWidget(self.boton_cargar)

        self.boton_exportar = QPushButton("Exportar resultados a Excel")
        self.boton_exportar.clicked.connect(self.exportar_excel)
        layout.addWidget(self.boton_exportar)

        # Filtros
        filtro_layout = QHBoxLayout()

        self.combo_empleados = QComboBox()
        self.combo_empleados.currentIndexChanged.connect(self.filtrar_empleado)

        self.boton_mostrar_todos = QPushButton("Mostrar todos")
        self.boton_mostrar_todos.clicked.connect(self.mostrar_todos)

        filtro_layout.addWidget(QLabel("Filtrar por empleado:"))
        filtro_layout.addWidget(self.combo_empleados)
        filtro_layout.addWidget(self.boton_mostrar_todos)

        layout.addLayout(filtro_layout)

        # Eliminar días / fechas
        eliminar_layout = QHBoxLayout()

        self.lista_dias = QListWidget()
        self.lista_dias.setSelectionMode(QListWidget.SelectionMode.MultiSelection)
        for d in ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]:
            self.lista_dias.addItem(d)

        eliminar_layout.addWidget(QLabel("Eliminar días de la semana:"))
        eliminar_layout.addWidget(self.lista_dias)

        self.input_fechas = QLineEdit()
        self.input_fechas.setPlaceholderText("Fechas a eliminar: dd/mm/yyyy, dd/mm/yyyy...")
        eliminar_layout.addWidget(self.input_fechas)

        self.boton_eliminar = QPushButton("Aplicar eliminación")
        self.boton_eliminar.clicked.connect(self.eliminar_dias_fechas)
        eliminar_layout.addWidget(self.boton_eliminar)

        layout.addLayout(eliminar_layout)

        self.tabla = QTableWidget()
        layout.addWidget(self.tabla)

        self.setLayout(layout)

    def parsear_fecha_hora(self, valor):
        from dateutil import parser

        if pd.isna(valor):
            return None

        s = str(valor).strip()
        if not s:
            return None

        # normalizar am/pm en español
        s = (
            s.replace("a. m.", "AM").replace("p. m.", "PM")
             .replace("a.m.", "AM").replace("p.m.", "PM")
             .replace("am", "AM").replace("pm", "PM")
             .replace("a m", "AM").replace("p m", "PM")
        )
        s = " ".join(s.split())

        try:
            return parser.parse(s, dayfirst=True)
        except Exception:
            try:
                return parser.parse(s, dayfirst=False)
            except Exception:
                return None

    def cargar_y_procesar_excel(self):
        archivo, _ = QFileDialog.getOpenFileName(
            self,
            "Seleccionar archivo Excel",
            "",
            "Archivos Excel (*.xlsx *.xls *.csv)"
        )
        if not archivo:
            return

        try:
            # 1) Leer robusto
            df_raw = cargar_excel_robusto(archivo)
            self._raw_rows = len(df_raw)

            # 2) Detectar columnas
            col_nombre, col_fechahora, _ = detectar_columnas_reloj(df_raw)

            if not col_nombre:
                raise Exception("No pude detectar la columna de NOMBRE.")
            if not col_fechahora:
                raise Exception("No pude detectar la columna de FECHA/HORA / MARCACIÓN.")

            df = df_raw.copy()
            df["Nombre"] = df[col_nombre].astype(str).str.strip()
            df["Marc."] = df[col_fechahora].apply(self.parsear_fecha_hora)

            # filas válidas
            df = df.dropna(subset=["Marc."]).copy()
            self._valid_rows = len(df)

            df["Fecha"] = df["Marc."].dt.date
            df = df.sort_values(["Nombre", "Marc."])

            # primer fichaje del día por persona
            df_primer = df.drop_duplicates(subset=["Nombre", "Fecha"], keep="first").copy()
            self._primeros_fichajes = len(df_primer)

            # hora
            df_primer["Hora"] = df_primer["Marc."].dt.time

            # solo antes de las 06:00
            df_pre6 = df_primer[df_primer["Hora"] < time(6, 0, 0)].copy()
            self._pre6_count = len(df_pre6)

            # calcular horas extra hasta las 06:00
            def calcular_extra(hora):
                limite = time(6, 0, 0)
                t1 = datetime.combine(datetime.today(), hora)
                t2 = datetime.combine(datetime.today(), limite)
                segundos = int((t2 - t1).total_seconds())
                if segundos < 0:
                    segundos = 0
                return str(timedelta(seconds=segundos))

            df_pre6["Hora Fichaje"] = df_pre6["Hora"].astype(str)
            df_pre6["Horas Extra"] = df_pre6["Hora"].apply(calcular_extra)
            df_pre6["Fecha"] = df_pre6["Fecha"].apply(lambda x: x.strftime("%d/%m/%Y"))

            # columnas finales
            self.df_base = df_pre6[["Nombre", "Fecha", "Hora Fichaje", "Horas Extra"]].reset_index(drop=True)
            self.df_procesado = self.df_base.copy()

            # combo empleados
            self.combo_empleados.clear()
            empleados = sorted(self.df_procesado["Nombre"].dropna().unique())
            self.combo_empleados.addItems(["-- Seleccione --"] + list(empleados))

            self.mostrar_en_tabla(self.df_procesado)
            self.actualizar_label_auditoria()

        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo procesar el archivo:\n{e}")

    def actualizar_label_auditoria(self):
        raw = getattr(self, "_raw_rows", 0)
        valid = getattr(self, "_valid_rows", 0)
        primeros = getattr(self, "_primeros_fichajes", 0)
        pre6 = getattr(self, "_pre6_count", 0)

        self.label.setText(
            f"✅ Aplicación 2 | Filas crudas: {raw} | Filas válidas: {valid} | "
            f"Primer fichaje/día: {primeros} | Antes de 06:00: {pre6}"
        )

    def mostrar_en_tabla(self, df=None):
        if df is None:
            df = self.df_procesado

        self.tabla.setRowCount(0)
        self.tabla.setColumnCount(0)

        if df is None or df.empty:
            QMessageBox.information(self, "Atención", "No hay fichajes antes de las 06:00 AM.")
            return

        self.tabla.setColumnCount(len(df.columns))
        self.tabla.setHorizontalHeaderLabels(df.columns.astype(str).tolist())

        for i, fila in df.iterrows():
            self.tabla.insertRow(i)
            for j, valor in enumerate(fila):
                item = QTableWidgetItem(str(valor))
                self.tabla.setItem(i, j, item)

            # resaltar entradas muy tempranas
            try:
                hora_str = str(fila.get("Hora Fichaje", ""))
                h, m, s = map(int, hora_str.split(":"))
                hora_obj = time(h, m, s)

                color = None
                if hora_obj < time(4, 0, 0):
                    color = QColor(255, 180, 180)
                elif hora_obj < time(5, 0, 0):
                    color = QColor(255, 220, 180)

                if color:
                    for j in range(len(fila)):
                        item = self.tabla.item(i, j)
                        if item:
                            item.setBackground(color)
            except Exception:
                pass

    def filtrar_empleado(self):
        if self.df_base is None or self.combo_empleados.currentIndex() <= 0:
            return

        nombre = self.combo_empleados.currentText()
        df_filtrado = self.df_base[self.df_base["Nombre"] == nombre].copy()

        td = pd.to_timedelta(df_filtrado["Horas Extra"].astype(str), errors="coerce").fillna(pd.Timedelta(0))
        total = str(td.sum())

        if not df_filtrado.empty:
            df_filtrado.loc[len(df_filtrado)] = {
                "Nombre": f"TOTAL DE HORAS EXTRAS DE {nombre}",
                "Fecha": "",
                "Hora Fichaje": "",
                "Horas Extra": total
            }

        self.df_procesado = df_filtrado
        self.mostrar_en_tabla(self.df_procesado)

    def mostrar_todos(self):
        if self.df_base is not None:
            self.df_procesado = self.df_base.copy()
            self.mostrar_en_tabla(self.df_procesado)
            self.actualizar_label_auditoria()

    def eliminar_dias_fechas(self):
        if self.df_base is None:
            return

        df = self.df_base.copy()

        dias_seleccionados = [item.text() for item in self.lista_dias.selectedItems()]
        if dias_seleccionados:
            dias_map = {
                "Lunes": 0,
                "Martes": 1,
                "Miércoles": 2,
                "Jueves": 3,
                "Viernes": 4,
                "Sábado": 5,
                "Domingo": 6,
            }
            wanted = {dias_map[d] for d in dias_seleccionados}

            idx_drop = []
            for i, row in df.iterrows():
                try:
                    fecha_obj = datetime.strptime(str(row["Fecha"]), "%d/%m/%Y").date()
                    if fecha_obj.weekday() in wanted:
                        idx_drop.append(i)
                except Exception:
                    pass

            df = df.drop(idx_drop)

        fechas_input = self.input_fechas.text().strip()
        if fechas_input:
            fechas = []
            for f in fechas_input.split(","):
                f = f.strip()
                if not f:
                    continue
                try:
                    fechas.append(datetime.strptime(f, "%d/%m/%Y").date())
                except Exception:
                    pass

            fechas_set = set(fechas)
            df = df[
                ~df["Fecha"].apply(
                    lambda x: datetime.strptime(str(x), "%d/%m/%Y").date() in fechas_set
                    if str(x).strip() else False
                )
            ]

        self.df_procesado = df.reset_index(drop=True)
        self.mostrar_en_tabla(self.df_procesado)

    def exportar_excel(self):
        if self.df_procesado is None or self.df_procesado.empty:
            QMessageBox.warning(self, "Atención", "No hay datos para exportar.")
            return

        archivo, _ = QFileDialog.getSaveFileName(
            self,
            "Guardar Excel",
            "",
            "Archivos Excel (*.xlsx)"
        )
        if not archivo:
            return

        try:
            df_export = self.df_procesado.copy()

            # hoja detalle con totales por persona
            df_final = pd.DataFrame()
            for nombre, grupo in df_export.groupby("Nombre"):
                df_final = pd.concat([df_final, grupo], ignore_index=True)

                td = pd.to_timedelta(grupo["Horas Extra"].astype(str), errors="coerce").fillna(pd.Timedelta(0))
                total_td = td.sum()

                if total_td > pd.Timedelta(0):
                    df_final = pd.concat(
                        [
                            df_final,
                            pd.DataFrame(
                                [
                                    {
                                        "Nombre": f"TOTAL DE HORAS EXTRAS DE {nombre}",
                                        "Fecha": "",
                                        "Hora Fichaje": "",
                                        "Horas Extra": str(total_td),
                                    }
                                ]
                            )
                        ],
                        ignore_index=True
                    )

            # hoja resumen
            tmp = df_export.copy()
            tmp["_td_extra"] = pd.to_timedelta(tmp["Horas Extra"].astype(str), errors="coerce").fillna(pd.Timedelta(0))

            resumen = (
                tmp.groupby("Nombre", as_index=False)
                .agg(
                    Dias=("Fecha", "count"),
                    Total_Horas_Extra=("_td_extra", "sum")
                )
            )
            resumen["Total Horas Extra"] = resumen["Total_Horas_Extra"].astype(str)
            resumen = resumen.drop(columns=["Total_Horas_Extra"])

            # hoja auditoría
            auditoria = pd.DataFrame(
                [
                    ["Filas crudas", getattr(self, "_raw_rows", 0)],
                    ["Filas válidas", getattr(self, "_valid_rows", 0)],
                    ["Primer fichaje por día", getattr(self, "_primeros_fichajes", 0)],
                    ["Fichajes antes de 06:00", getattr(self, "_pre6_count", 0)],
                ],
                columns=["Métrica", "Valor"]
            )

            with pd.ExcelWriter(archivo, engine="openpyxl") as writer:
                df_final.to_excel(writer, index=False, sheet_name="Detalle_Pre6")
                resumen.to_excel(writer, index=False, sheet_name="Resumen")
                auditoria.to_excel(writer, index=False, sheet_name="Auditoria")

                ws = writer.sheets["Detalle_Pre6"]
                amarillo = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

                for i, row in enumerate(df_final.itertuples(), start=2):
                    if "TOTAL DE HORAS EXTRAS" in str(row.Nombre):
                        for j in range(1, len(df_final.columns) + 1):
                            ws.cell(row=i, column=j).fill = amarillo

            QMessageBox.information(self, "Éxito", "Archivo exportado correctamente.")

        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo exportar el archivo:\n{e}")
# ============================================
# 🌟 VENTANA PRINCIPAL
# ============================================
class VentanaPrincipal(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("📊 Sistema de Control de Fichajes")
        self.resize(1200, 700)
        layout = QVBoxLayout()

        self.tabs = QTabWidget()
        self.tabs.addTab(Aplicacion1(), "Fichajes Estándar")
        self.tabs.addTab(Aplicacion2(), "Horas antes de 06 AM")
        self.tabs.addTab(AplicacionPresentismo(), "Presentismo")

        layout.addWidget(self.tabs)
        self.setLayout(layout)


# ============================================
# 🔹 EJECUCIÓN
# ============================================
if __name__ == "__main__":
    app = QApplication(sys.argv)
    ventana = VentanaPrincipal()
    ventana.show()
    sys.exit(app.exec())
