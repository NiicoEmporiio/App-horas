import sys
import pandas as pd
from PyQt6.QtWidgets import (
    QApplication, QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QLabel,
    QTableWidget, QTableWidgetItem, QMessageBox, QFileDialog,
    QTabWidget, QComboBox, QLineEdit, QListWidget
)
from PyQt6.QtCore import Qt
from datetime import datetime, timedelta
from PyQt6.QtGui import QColor
from openpyxl.styles import PatternFill

# ============================================
# 🟢 CLASE BASE GENERAL PARA LEER ARCHIVOS EXCEL
# ============================================
class PestañaExcel(QWidget):
    """Clase base para pestañas simples de carga de Excel"""
    def __init__(self, titulo):
        super().__init__()
        self.titulo = titulo
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()
        self.label = QLabel(f"📂 {self.titulo}")
        self.label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self.label)

        self.boton_cargar = QPushButton("Seleccionar archivo Excel")
        self.boton_cargar.clicked.connect(self.cargar_excel)
        layout.addWidget(self.boton_cargar)

        self.tabla = QTableWidget()
        layout.addWidget(self.tabla)

        self.setLayout(layout)

    def cargar_excel(self):
        archivo, _ = QFileDialog.getOpenFileName(
            self, "Seleccionar archivo Excel", "", "Archivos Excel (*.xlsx *.xls)"
        )
        if archivo:
            try:
                try:
                    df = pd.read_excel(archivo, engine="openpyxl")
                except Exception:
                    df = pd.read_excel(archivo, engine="xlrd")
                self.mostrar_en_tabla(df)
            except Exception as e:
                QMessageBox.critical(self, "Error", f"No se pudo leer el archivo:\n{e}")

    def mostrar_en_tabla(self, df):
        self.tabla.setRowCount(0)
        self.tabla.setColumnCount(0)
        if df.empty:
            QMessageBox.warning(self, "Atención", "El archivo Excel está vacío.")
            return
        self.tabla.setColumnCount(len(df.columns))
        self.tabla.setHorizontalHeaderLabels(df.columns.astype(str).tolist())
        for i, fila in df.iterrows():
            self.tabla.insertRow(i)
            for j, valor in enumerate(fila):
                item = QTableWidgetItem(str(valor))
                self.tabla.setItem(i, j, item)
        self.label.setText(f"✅ {self.titulo} - Archivo cargado correctamente ({len(df)} filas)")

# ============================================
# 🔵 APLICACIÓN 1 - CONTROL DE FICHAJE Y HORAS EXTRAS
# ============================================
class Aplicacion1(QWidget):
    """Procesa fichajes: limpia, agrupa, calcula horas trabajadas, horas extras,
       permite filtrar por empleado y eliminar días/fechas"""
    def __init__(self):
        super().__init__()
        self.df_procesado = None
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()
        self.label = QLabel("📊 Aplicación 1: Control de fichaje y horas extras")
        self.label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self.label)

        self.boton_cargar = QPushButton("Seleccionar archivo Excel de fichajes")
        self.boton_cargar.clicked.connect(self.cargar_y_procesar_excel)
        layout.addWidget(self.boton_cargar)

        self.boton_exportar = QPushButton("Exportar resultados a Excel")
        self.boton_exportar.clicked.connect(self.exportar_excel)
        layout.addWidget(self.boton_exportar)

        # Filtro por empleado
        filtro_layout = QHBoxLayout()
        self.combo_empleados = QComboBox()
        self.combo_empleados.currentIndexChanged.connect(self.filtrar_empleado)
        self.boton_mostrar_todos = QPushButton("Mostrar todos")
        self.boton_mostrar_todos.clicked.connect(self.mostrar_todos)
        filtro_layout.addWidget(QLabel("Filtrar por empleado:"))
        filtro_layout.addWidget(self.combo_empleados)
        filtro_layout.addWidget(self.boton_mostrar_todos)
        layout.addLayout(filtro_layout)

        # Eliminación por día o fecha
        eliminar_layout = QHBoxLayout()
        self.lista_dias = QListWidget()
        self.lista_dias.setSelectionMode(QListWidget.SelectionMode.MultiSelection)
        dias_semana = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
        for d in dias_semana:
            self.lista_dias.addItem(d)
        eliminar_layout.addWidget(QLabel("Eliminar días de la semana:"))
        eliminar_layout.addWidget(self.lista_dias)
        self.input_fechas = QLineEdit()
        self.input_fechas.setPlaceholderText("Fechas a eliminar: dd/mm/yyyy, dd/mm/yyyy...")
        eliminar_layout.addWidget(self.input_fechas)
        self.boton_eliminar = QPushButton("Aplicar eliminación")
        self.boton_eliminar.clicked.connect(self.eliminar_dias_fechas)
        eliminar_layout.addWidget(self.boton_eliminar)
        layout.addLayout(eliminar_layout)

        self.tabla = QTableWidget()
        layout.addWidget(self.tabla)

        self.setLayout(layout)

    def cargar_y_procesar_excel(self):
        archivo, _ = QFileDialog.getOpenFileName(
            self, "Seleccionar archivo Excel", "", "Archivos Excel (*.xlsx *.xls)"
        )
        if archivo:
            try:
                try:
                    df = pd.read_excel(archivo, engine="openpyxl")
                except Exception:
                    df = pd.read_excel(archivo, engine="xlrd")
                self.df_procesado = self.procesar_fichajes(df)
                self.combo_empleados.clear()
                empleados = sorted(self.df_procesado["Nombre"].unique())
                self.combo_empleados.addItems(["-- Seleccione --"] + empleados)
                self.mostrar_en_tabla(self.df_procesado)
            except Exception as e:
                QMessageBox.critical(self, "Error", f"No se pudo procesar el archivo:\n{e}")

    def procesar_fichajes(self, df):
        columnas = ["Nombre", "Fecha/Hora"]
        for col in columnas:
            if col not in df.columns:
                raise Exception(f"Falta la columna requerida: {col}")
        df["Fecha/Hora"] = pd.to_datetime(df["Fecha/Hora"], errors="coerce")
        df["Fecha"] = df["Fecha/Hora"].dt.date
        df["Hora"] = df["Fecha/Hora"].dt.time
        df = df.dropna(subset=["Fecha/Hora"])
        resultado = []
        for (nombre, fecha), grupo in df.groupby(["Nombre", "Fecha"]):
            grupo_ordenado = grupo.sort_values("Fecha/Hora")
            hora_entrada = grupo_ordenado.iloc[0]["Hora"] if not grupo_ordenado.empty else None
            hora_salida = grupo_ordenado.iloc[-1]["Hora"] if len(grupo_ordenado) > 1 else None
            if len(grupo_ordenado) == 0:
                entrada = salida = total_horas = horas_extra = ""
                estado = "Sin fichaje"
            elif len(grupo_ordenado) == 1:
                if grupo_ordenado.iloc[0]["Hora"].hour < 12:
                    entrada = hora_entrada.strftime("%H:%M:%S")
                    salida = total_horas = horas_extra = ""
                    estado = "⚠️ Falta salida"
                else:
                    entrada = total_horas = horas_extra = ""
                    salida = hora_entrada.strftime("%H:%M:%S")
                    estado = "⚠️ Falta entrada"
            else:
                entrada = hora_entrada.strftime("%H:%M:%S") if hora_entrada else ""
                salida = hora_salida.strftime("%H:%M:%S") if hora_salida else ""
                estado = "✅ Completo"
                tiempo_entrada = datetime.combine(datetime.today(), hora_entrada)
                tiempo_salida = datetime.combine(datetime.today(), hora_salida)
                total = tiempo_salida - tiempo_entrada
                total_horas = str(total).split('.')[0]
                limite = timedelta(hours=8)
                extra = total - limite if total > limite else timedelta(0)
                horas_extra = str(extra).split('.')[0]
            resultado.append({
                "Nombre": nombre,
                "Fecha": fecha.strftime("%d/%m/%Y"),
                "Entrada": entrada,
                "Salida": salida,
                "Estado": estado,
                "Total Horas": total_horas,
                "Horas Extra": horas_extra
            })
        return pd.DataFrame(resultado)

    def mostrar_en_tabla(self, df):
        self.tabla.setRowCount(0)
        self.tabla.setColumnCount(0)
        if df.empty:
            QMessageBox.warning(self, "Atención", "No se encontraron fichajes válidos.")
            return
        self.tabla.setColumnCount(len(df.columns))
        self.tabla.setHorizontalHeaderLabels(df.columns.astype(str).tolist())
        for i, fila in df.iterrows():
            self.tabla.insertRow(i)
            for j, valor in enumerate(fila):
                item = QTableWidgetItem(str(valor))
                self.tabla.setItem(i, j, item)
            color = None
            if fila.get("Estado") == "✅ Completo":
                color = QColor(200, 255, 200)
            elif fila.get("Estado") and "⚠️" in fila.get("Estado"):
                color = QColor(255, 255, 150)
            elif fila.get("Estado") == "Sin fichaje":
                color = QColor(255, 200, 200)
            if color:
                for j in range(len(fila)):
                    item = self.tabla.item(i, j)
                    if item:
                        item.setBackground(color)
        self.label.setText(f"✅ Fichajes procesados correctamente ({len(df)} registros)")

    def filtrar_empleado(self):
        if self.df_procesado is not None and self.combo_empleados.currentIndex() > 0:
            nombre = self.combo_empleados.currentText()
            df_filtrado = self.df_procesado[self.df_procesado["Nombre"] == nombre].copy()
            total_extra_segundos = 0
            for val in df_filtrado["Horas Extra"]:
                if val and val != "0:00:00":
                    h, m, s = map(int, val.split(":"))
                    total_extra_segundos += h*3600 + m*60 + s
            total_extra = str(timedelta(seconds=total_extra_segundos))
            if not df_filtrado.empty:
                df_filtrado.loc[len(df_filtrado)] = {
                    "Nombre": f"TOTAL DE HORAS EXTRAS DE {nombre}",
                    "Fecha": "",
                    "Entrada": "",
                    "Salida": "",
                    "Estado": "",
                    "Total Horas": "",
                    "Horas Extra": total_extra
                }
            self.mostrar_en_tabla(df_filtrado)

    def mostrar_todos(self):
        if self.df_procesado is not None:
            self.mostrar_en_tabla(self.df_procesado)

    def eliminar_dias_fechas(self):
        if self.df_procesado is None:
            return
        df = self.df_procesado.copy()
        dias_seleccionados = [item.text() for item in self.lista_dias.selectedItems()]
        if dias_seleccionados:
            dias_map = {"Lunes":0,"Martes":1,"Miércoles":2,"Jueves":3,"Viernes":4,"Sábado":5,"Domingo":6}
            indices_a_eliminar = []
            for i, row in df.iterrows():
                fecha_obj = datetime.strptime(row["Fecha"], "%d/%m/%Y").date()
                if fecha_obj.weekday() in [dias_map[d] for d in dias_seleccionados]:
                    indices_a_eliminar.append(i)
            df = df.drop(indices_a_eliminar)
        fechas_input = self.input_fechas.text().strip()
        if fechas_input:
            fechas = [datetime.strptime(f.strip(), "%d/%m/%Y").date() for f in fechas_input.split(",") if f.strip()]
            df = df[~df["Fecha"].apply(lambda x: datetime.strptime(x,"%d/%m/%Y").date() in fechas)]
        self.df_procesado = df.reset_index(drop=True)
        self.mostrar_en_tabla(self.df_procesado)

    def exportar_excel(self):
        if self.df_procesado is None or self.df_procesado.empty:
            QMessageBox.warning(self, "Atención", "No hay datos para exportar.")
            return
        archivo, _ = QFileDialog.getSaveFileName(self, "Guardar Excel", "", "Archivos Excel (*.xlsx)")
        if not archivo:
            return
        try:
            with pd.ExcelWriter(archivo, engine="openpyxl") as writer:
                df_export = self.df_procesado.copy()
                df_final = pd.DataFrame()
                for nombre, grupo in df_export.groupby("Nombre"):
                    df_final = pd.concat([df_final, grupo], ignore_index=True)
                    total_segundos = 0
                    for val in grupo["Horas Extra"]:
                        if val and val != "0:00:00":
                            h, m, s = map(int, val.split(":"))
                            total_segundos += h*3600 + m*60 + s
                    if total_segundos > 0:
                        df_final = pd.concat([df_final, pd.DataFrame([{
                            "Nombre": f"TOTAL HORAS EXTRAS DE {nombre}",
                            "Fecha": "",
                            "Entrada": "",
                            "Salida": "",
                            "Estado": "",
                            "Total Horas": "",
                            "Horas Extra": str(timedelta(seconds=total_segundos))
                        }])], ignore_index=True)
                df_final.to_excel(writer, index=False, sheet_name="Datos Procesados")
                ws = writer.sheets["Datos Procesados"]
                amarillo = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
                for i, row in enumerate(df_final.itertuples(), start=2):
                    if "TOTAL HORAS EXTRAS" in str(row.Nombre):
                        for j in range(1, len(df_final.columns)+1):
                            ws.cell(row=i, column=j).fill = amarillo
                resumen = df_export.groupby("Nombre")["Horas Extra"].apply(
                    lambda x: str(timedelta(
                        seconds=sum(
                            int(t.split(":")[0])*3600 + int(t.split(":")[1])*60 + int(t.split(":")[2])
                            for t in x if t and t != "0:00:00"
                        )
                    ))
                ).reset_index()
                resumen.rename(columns={"Horas Extra": "Total Horas Extra"}, inplace=True)
                resumen.to_excel(writer, index=False, sheet_name="Total Horas Extra")
            QMessageBox.information(self, "Éxito", "Archivo exportado correctamente.")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo exportar el archivo:\n{e}")

# ============================================
# 🟢 APLICACIÓN 2 - HORAS PREVIAS A LAS 6 AM
# ============================================
class Aplicacion2(QWidget):
    """Procesa fichajes antes de las 06:00 AM"""
    def __init__(self):
        super().__init__()
        self.df_procesado = None
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()
        self.label = QLabel("📊 Aplicación 2: Horas previas a las 6 AM")
        self.label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self.label)
        self.boton_cargar = QPushButton("Seleccionar archivo Excel")
        self.boton_cargar.clicked.connect(self.cargar_y_procesar_excel)
        layout.addWidget(self.boton_cargar)
        self.boton_exportar = QPushButton("Exportar resultados a Excel")
        self.boton_exportar.clicked.connect(self.exportar_excel)
        layout.addWidget(self.boton_exportar)
        self.tabla = QTableWidget()
        layout.addWidget(self.tabla)
        self.setLayout(layout)

    def cargar_y_procesar_excel(self):
        archivo, _ = QFileDialog.getOpenFileName(self, "Seleccionar archivo Excel", "", "Archivos Excel (*.xlsx *.xls)")
        if archivo:
            try:
                try:
                    df = pd.read_excel(archivo, engine="openpyxl")
                except Exception:
                    df = pd.read_excel(archivo, engine="xlrd")
                self.df_procesado = self.procesar_fichajes_pre6(df)
                self.mostrar_en_tabla(self.df_procesado)
            except Exception as e:
                QMessageBox.critical(self, "Error", f"No se pudo procesar el archivo:\n{e}")

    def procesar_fichajes_pre6(self, df):
        columnas = ["Nombre","Marc."]
        for col in columnas:
            if col not in df.columns:
                raise Exception(f"Falta columna requerida: {col}")
        df["Marc."] = pd.to_datetime(df["Marc."], errors="coerce")
        df = df.dropna(subset=["Marc."])
        df["Fecha"] = df["Marc."].dt.date
        df["Hora"] = df["Marc."].dt.time
        df = df[df["Marc."].dt.hour < 6]
        df = df.sort_values(["Nombre","Fecha","Marc."])
        df = df.drop_duplicates(subset=["Nombre","Fecha"], keep="first")
        resultado = []
        limite = datetime.strptime("06:00:00","%H:%M:%S").time()
        for i,fila in df.iterrows():
            entrada = fila["Hora"]
            tiempo_extra = datetime.combine(datetime.today(), limite) - datetime.combine(datetime.today(), entrada)
            horas_extra = str(tiempo_extra).split('.')[0]
            resultado.append({
                "Nombre":fila["Nombre"],
                "Fecha":fila["Fecha"].strftime("%d/%m/%Y"),
                "Entrada":entrada.strftime("%H:%M:%S"),
                "Horas Extra":horas_extra
            })
        return pd.DataFrame(resultado)

    def mostrar_en_tabla(self, df):
        self.tabla.setRowCount(0)
        self.tabla.setColumnCount(0)
        if df.empty:
            QMessageBox.warning(self,"Atención","No hay fichajes previos a las 06:00 AM.")
            return
        self.tabla.setColumnCount(len(df.columns))
        self.tabla.setHorizontalHeaderLabels(df.columns.astype(str).tolist())
        for i,fila in df.iterrows():
            self.tabla.insertRow(i)
            for j,valor in enumerate(fila):
                item = QTableWidgetItem(str(valor))
                self.tabla.setItem(i,j,item)
        self.label.setText(f"✅ Fichajes pre6 procesados ({len(df)} registros)")

    def exportar_excel(self):
        if self.df_procesado is None or self.df_procesado.empty:
            QMessageBox.warning(self,"Atención","No hay datos para exportar.")
            return
        archivo, _ = QFileDialog.getSaveFileName(self,"Guardar Excel","","Archivos Excel (*.xlsx)")
        if not archivo:
            return
        try:
            with pd.ExcelWriter(archivo, engine="openpyxl") as writer:
                df_export = self.df_procesado.copy()
                df_final = pd.DataFrame()
                for nombre,grupo in df_export.groupby("Nombre"):
                    df_final = pd.concat([df_final,grupo],ignore_index=True)
                    total_segundos = sum(int(t.split(":")[0])*3600+int(t.split(":")[1])*60+int(t.split(":")[2]) for t in grupo["Horas Extra"])
                    df_final = pd.concat([df_final,pd.DataFrame([{
                        "Nombre":f"TOTAL HORAS EXTRAS DE {nombre}",
                        "Fecha":"",
                        "Entrada":"",
                        "Horas Extra":str(timedelta(seconds=total_segundos))
                    }])],ignore_index=True)
                df_final.to_excel(writer,index=False,sheet_name="Datos pre6")
                ws = writer.sheets["Datos pre6"]
                amarillo = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
                for i,row in enumerate(df_final.itertuples(),start=2):
                    if "TOTAL HORAS EXTRAS" in str(row.Nombre):
                        for j in range(1,len(df_final.columns)+1):
                            ws.cell(row=i,column=j).fill = amarillo
                resumen = df_export.groupby("Nombre")["Horas Extra"].apply(
                    lambda x: str(timedelta(
                        seconds=sum(int(t.split(":")[0])*3600+int(t.split(":")[1])*60+int(t.split(":")[2]) for t in x)
                    ))
                ).reset_index()
                resumen.rename(columns={"Horas Extra":"Total Horas Extra"}, inplace=True)
                resumen.to_excel(writer,index=False,sheet_name="Resumen Total")
            QMessageBox.information(self,"Éxito","Archivo exportado correctamente.")
        except Exception as e:
            QMessageBox.critical(self,"Error",f"No se pudo exportar el archivo:\n{e}")

# ============================================
# 🟢 APLICACIÓN 3 - SIMPLE CARGA DE EXCEL (PRESENTISMO)
# ============================================
class Aplicacion3(PestañaExcel):
    def __init__(self):
        super().__init__("Aplicación 3: Presentismo base")

# ============================================
# 🌟 VENTANA PRINCIPAL
# ============================================
class VentanaPrincipal(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Sistema de Fichajes y Horas Extras")
        self.resize(1200,700)
        layout = QVBoxLayout()
        self.tabs = QTabWidget()
        self.tabs.addTab(Aplicacion1(),"📋 Fichajes Generales")
        self.tabs.addTab(Aplicacion2(),"⏰ Horas pre6")
        self.tabs.addTab(Aplicacion3(),"🟢 Presentismo Base")
        layout.addWidget(self.tabs)
        self.setLayout(layout)

# ============================================
# 🚀 EJECUCIÓN
# ============================================
if __name__ == "__main__":
    app = QApplication(sys.argv)
    ventana = VentanaPrincipal()
    ventana.show()
    sys.exit(app.exec())
