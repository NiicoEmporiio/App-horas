import sys
import pandas as pd
import re

from PyQt6.QtWidgets import (
    QApplication, QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QLabel,
    QTableWidget, QTableWidgetItem, QMessageBox, QFileDialog,
    QTabWidget, QComboBox, QLineEdit, QListWidget
)
from PyQt6.QtCore import Qt
from datetime import datetime, timedelta, time
from PyQt6.QtGui import QColor
from openpyxl.styles import PatternFill
from aplicacion_presentismo import AplicacionPresentismo


# ============================================
# 🟢 CLASE BASE GENERAL PARA LEER ARCHIVOS EXCEL
# ============================================


def _norm_col(s: str) -> str:
    s = str(s).strip().lower()
    s = s.replace("\n", " ").replace("\r", " ")
    s = re.sub(r"\s+", " ", s)
    return s

def _find_col(df, candidates_exact=(), candidates_contains=()):
    cols = list(df.columns)
    norm_map = {c: _norm_col(c) for c in cols}

    exact_norm = {_norm_col(x) for x in candidates_exact}
    for c, nc in norm_map.items():
        if nc in exact_norm:
            return c

    contains_norm = [_norm_col(x) for x in candidates_contains]
    for c, nc in norm_map.items():
        for pat in contains_norm:
            if pat and pat in nc:
                return c

    return None


def cargar_excel_robusto(ruta: str) -> pd.DataFrame:
    """Lee xlsx/xls/csv de forma robusta y devuelve un DataFrame."""
    ruta_lower = str(ruta).lower().strip()

    # CSV
    if ruta_lower.endswith(".csv"):
        try:
            return pd.read_csv(ruta, encoding="utf-8")
        except Exception:
            return pd.read_csv(ruta, encoding="latin-1")

    # Excel
    try:
        return pd.read_excel(ruta, engine="openpyxl")  # xlsx
    except Exception:
        pass

    try:
        return pd.read_excel(ruta, engine="xlrd")      # xls
    except Exception:
        pass

    # fallback
    return pd.read_excel(ruta)

def detectar_columnas_reloj(df):
    col_nombre = _find_col(
        df,
        candidates_exact=("Nombre", "Empleado", "Personal", "Funcionario"),
        candidates_contains=("nombre", "emple", "funcion", "personal"),
    )

    col_fechahora = _find_col(
        df,
        candidates_exact=("Fecha/Hora", "Fecha Hora", "Fecha y Hora", "Marc.", "Marcación", "Marcacion", "Timestamp"),
        candidates_contains=("fecha/hora", "fecha hora", "marc", "timestamp", "fecha y hora"),
    )

    col_tipo = _find_col(
        df,
        candidates_exact=("Tipo de registro", "Tipo", "Evento", "Movimiento"),
        candidates_contains=("tipo de registro", "tipo", "evento", "movim"),
    )

    return col_nombre, col_fechahora, col_tipo
class PestañaExcel(QWidget):
    """Clase base para pestañas simples de carga de Excel"""
    def __init__(self, titulo):
        super().__init__()
        self.titulo = titulo
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()
        self.label = QLabel(f"📂 {self.titulo}")
        self.label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self.label)

        self.boton_cargar = QPushButton("Seleccionar archivo Excel")
        self.boton_cargar.clicked.connect(self.cargar_excel)
        layout.addWidget(self.boton_cargar)

        self.tabla = QTableWidget()
        layout.addWidget(self.tabla)

        self.setLayout(layout)
        

    def cargar_excel(self):
        archivo, _ = QFileDialog.getOpenFileName(
            self, "Seleccionar archivo Excel", "", "Archivos Excel (*.xlsx *.xls)"
        )
        if archivo:
            try:
                try:
                    df = pd.read_excel(archivo, engine="openpyxl")
                except Exception:
                    df = pd.read_excel(archivo, engine="xlrd")
                self.mostrar_en_tabla(df)
            except Exception as e:
                QMessageBox.critical(self, "Error", f"No se pudo leer el archivo:\n{e}")

    def mostrar_en_tabla(self, df):
        self.tabla.setRowCount(0)
        self.tabla.setColumnCount(0)
        if df.empty:
            QMessageBox.warning(self, "Atención", "El archivo Excel está vacío.")
            return
        self.tabla.setColumnCount(len(df.columns))
        self.tabla.setHorizontalHeaderLabels(df.columns.astype(str).tolist())
        for i, fila in df.iterrows():
            self.tabla.insertRow(i)
            for j, valor in enumerate(fila):
                item = QTableWidgetItem(str(valor))
                self.tabla.setItem(i, j, item)
        self.label.setText(f"✅ {self.titulo} - Archivo cargado correctamente ({len(df)} filas)")

# ============================================
# 🔵 APLICACIÓN 1 - CONTROL DE FICHAJE Y HORAS EXTRAS
# ============================================
class Aplicacion1(QWidget):
    """Procesa fichajes: limpia, agrupa, calcula horas trabajadas, horas extras,
       permite filtrar por empleado y eliminar días/fechas"""
    def __init__(self):
        super().__init__()
        self.df_procesado = None
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()
        self.label = QLabel("📊 Aplicación 1: Control de fichaje y horas extras")
        self.label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self.label)

        # Botón cargar archivo
        self.boton_cargar = QPushButton("Seleccionar archivo Excel de fichajes")
        self.boton_cargar.clicked.connect(self.cargar_y_procesar_excel)
        layout.addWidget(self.boton_cargar)

        # Exportar Excel
        self.boton_exportar = QPushButton("Exportar resultados a Excel")
        self.boton_exportar.clicked.connect(self.exportar_excel)
        layout.addWidget(self.boton_exportar)

        # Filtro por empleado (ComboBox)
        filtro_layout = QHBoxLayout()
        self.combo_empleados = QComboBox()
        self.combo_empleados.currentIndexChanged.connect(self.filtrar_empleado)
        self.boton_mostrar_todos = QPushButton("Mostrar todos")
        self.boton_mostrar_todos.clicked.connect(self.mostrar_todos)
        filtro_layout.addWidget(QLabel("Filtrar por empleado:"))
        filtro_layout.addWidget(self.combo_empleados)
        filtro_layout.addWidget(self.boton_mostrar_todos)
        layout.addLayout(filtro_layout)

        # Eliminación por día o fecha
        eliminar_layout = QHBoxLayout()
        self.lista_dias = QListWidget()
        self.lista_dias.setSelectionMode(QListWidget.SelectionMode.MultiSelection)
        dias_semana = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
        for d in dias_semana:
            self.lista_dias.addItem(d)
        eliminar_layout.addWidget(QLabel("Eliminar días de la semana:"))
        eliminar_layout.addWidget(self.lista_dias)
        self.input_fechas = QLineEdit()
        self.input_fechas.setPlaceholderText("Fechas a eliminar: dd/mm/yyyy, dd/mm/yyyy...")
        eliminar_layout.addWidget(self.input_fechas)
        self.boton_eliminar = QPushButton("Aplicar eliminación")
        self.boton_eliminar.clicked.connect(self.eliminar_dias_fechas)
        eliminar_layout.addWidget(self.boton_eliminar)
        layout.addLayout(eliminar_layout)

        # Tabla para mostrar datos
        self.tabla = QTableWidget()
        layout.addWidget(self.tabla)

        self.setLayout(layout)

    # --------------------------
    # Cargar y procesar Excel
    # --------------------------
def cargar_y_procesar_excel(self):
    archivo, _ = QFileDialog.getOpenFileName(
        self,
        "Seleccionar archivo",
        "",
        "Archivos (*.xlsx *.xls *.csv)",
    )
    if not archivo:
        return

    try:
        # 1) Leer archivo robusto
        df_raw = cargar_excel_robusto(archivo)
        self._raw_rows = len(df_raw)

        # 2) Procesar
        self.df_procesado = self.procesar_fichajes(df_raw)

        # 3) Combo empleados
        self.combo_empleados.clear()
        empleados = sorted(self.df_procesado["Nombre"].dropna().unique())
        self.combo_empleados.addItems(["-- Seleccione --"] + list(empleados))

        # 4) Mostrar tabla
        self.mostrar_en_tabla(self.df_procesado)

        # 5) Auditoría en label
        raw = getattr(self, "_raw_rows", 0)
        valid = getattr(self, "_valid_rows", 0)
        completos = getattr(self, "_dias_completos", 0)
        incompletos = getattr(self, "_dias_incompletos", 0)
        sinfichaje = getattr(self, "_dias_sin", 0)

        self.label.setText(
            f"✅ Procesado | Filas crudas: {raw} | Filas válidas: {valid} | "
            f"Días completos: {completos} | Incompletos: {incompletos} | Sin fichaje: {sinfichaje}"
        )

    except Exception as e:
        QMessageBox.critical(self, "Error", f"No se pudo procesar el archivo:\n{e}")

    # --------------------------
    # Procesamiento de fichajes y cálculo de horas
    # --------------------------
    def procesar_fichajes(self, df):
        col_nombre, col_fechahora, col_tipo = detectar_columnas_reloj(df)

        if not col_nombre:
            raise Exception("No pude detectar la columna de NOMBRE (ej: 'Nombre' / 'Empleado').")
        if not col_fechahora:
            raise Exception("No pude detectar la columna de FECHA/HORA (ej: 'Fecha/Hora' / 'Marc.' / 'Timestamp').")

        df = df.copy()

        df["Nombre"] = df[col_nombre].astype(str).str.strip()
        df["Fecha/Hora"] = pd.to_datetime(df[col_fechahora], errors="coerce")
        df = df.dropna(subset=["Fecha/Hora"])
        self._valid_rows = len(df)
        
        
        df["Fecha"] = df["Fecha/Hora"].dt.date
        df["Hora"] = df["Fecha/Hora"].dt.time

        if col_tipo and col_tipo in df.columns:
            df["Tipo de registro"] = df[col_tipo].astype(str).str.strip().str.lower()

        resultado = []
        for (nombre, fecha), grupo in df.groupby(["Nombre", "Fecha"]):
            grupo_ordenado = grupo.sort_values("Fecha/Hora")
            hora_entrada = grupo_ordenado.iloc[0]["Hora"] if not grupo_ordenado.empty else None
            hora_salida = grupo_ordenado.iloc[-1]["Hora"] if len(grupo_ordenado) > 1 else None

            if len(grupo_ordenado) == 0:
                entrada = salida = total_horas = horas_extra = ""
                estado = "Sin fichaje"

            elif len(grupo_ordenado) == 1:
                if grupo_ordenado.iloc[0]["Hora"].hour < 12:
                    entrada = hora_entrada.strftime("%H:%M:%S")
                    salida = total_horas = horas_extra = ""
                    estado = "⚠️ Falta salida"
                else:
                    entrada = total_horas = horas_extra = ""
                    salida = hora_entrada.strftime("%H:%M:%S")
                    estado = "⚠️ Falta entrada"

            else:
                entrada = hora_entrada.strftime("%H:%M:%S") if hora_entrada else ""
                salida = hora_salida.strftime("%H:%M:%S") if hora_salida else ""
                estado = "✅ Completo"

                tiempo_entrada = datetime.combine(datetime.today(), hora_entrada)
                tiempo_salida = datetime.combine(datetime.today(), hora_salida)
                total = tiempo_salida - tiempo_entrada

                limite = timedelta(hours=8)
                extra = total - limite if total > limite else timedelta(0)

                total_horas = str(total).split(".")[0]
                horas_extra = str(extra).split(".")[0]

            resultado.append(
                {
                    "Nombre": nombre,
                    "Fecha": fecha.strftime("%d/%m/%Y"),
                    "Entrada": entrada,
                    "Salida": salida,
                    "Estado": estado,
                    "Total Horas": total_horas,
                    "Horas Extra": horas_extra,
                }
            )
        
        return pd.DataFrame(resultado)
    
    # --------------------------
    # Mostrar tabla con colores por estado
    # --------------------------
    def mostrar_en_tabla(self, df):
        self.tabla.setRowCount(0)
        self.tabla.setColumnCount(0)
        if df.empty:
            QMessageBox.warning(self, "Atención", "No se encontraron fichajes válidos.")
            return

        self.tabla.setColumnCount(len(df.columns))
        self.tabla.setHorizontalHeaderLabels(df.columns.astype(str).tolist())

        for i, fila in df.iterrows():
            self.tabla.insertRow(i)
            for j, valor in enumerate(fila):
                item = QTableWidgetItem(str(valor))
                self.tabla.setItem(i, j, item)

            # Colorear según estado
            color = None
            if fila.get("Estado") == "✅ Completo":
                color = QColor(200, 255, 200)
            elif fila.get("Estado") and "⚠️" in fila.get("Estado"):
                color = QColor(255, 255, 150)
            elif fila.get("Estado") == "Sin fichaje":
                color = QColor(255, 200, 200)
            if color:
                for j in range(len(fila)):
                    item = self.tabla.item(i, j)
                    if item:
                        item.setBackground(color)

        self.label.setText(f"✅ Fichajes procesados correctamente ({len(df)} registros)")

    # --------------------------
    # Filtro por empleado
    # --------------------------
    def filtrar_empleado(self):
        if self.df_procesado is not None and self.combo_empleados.currentIndex() > 0:
            nombre = self.combo_empleados.currentText()
            df_filtrado = self.df_procesado[self.df_procesado["Nombre"] == nombre].copy()

            # Calcular total horas extras al final
            total_extra_segundos = 0
            for val in df_filtrado["Horas Extra"]:
                if val and val != "0:00:00":
                    h, m, s = map(int, val.split(":"))
                    total_extra_segundos += h*3600 + m*60 + s
            total_extra = str(timedelta(seconds=total_extra_segundos))
            if not df_filtrado.empty:
                df_filtrado.loc[len(df_filtrado)] = {
                    "Nombre": f"TOTAL DE HORAS EXTRAS DE {nombre}",
                    "Fecha": "",
                    "Entrada": "",
                    "Salida": "",
                    "Estado": "",
                    "Total Horas": "",
                    "Horas Extra": total_extra
                }
            self.mostrar_en_tabla(df_filtrado)

    def mostrar_todos(self):
        if self.df_procesado is not None:
            self.mostrar_en_tabla(self.df_procesado)

    # --------------------------
    # Eliminar días por día de semana o fechas específicas
    # --------------------------
    def eliminar_dias_fechas(self):
        if self.df_procesado is None:
            return

        df = self.df_procesado.copy()

        # Eliminar días de la semana seleccionados
        dias_seleccionados = [item.text() for item in self.lista_dias.selectedItems()]
        if dias_seleccionados:
            dias_map = {"Lunes":0,"Martes":1,"Miércoles":2,"Jueves":3,"Viernes":4,"Sábado":5,"Domingo":6}
            indices_a_eliminar = []
            for i, row in df.iterrows():
                fecha_obj = datetime.strptime(row["Fecha"], "%d/%m/%Y").date()
                if fecha_obj.weekday() in [dias_map[d] for d in dias_seleccionados]:
                    indices_a_eliminar.append(i)
            df = df.drop(indices_a_eliminar)

        # Eliminar fechas específicas
        fechas_input = self.input_fechas.text().strip()
        if fechas_input:
            fechas = [datetime.strptime(f.strip(), "%d/%m/%Y").date() for f in fechas_input.split(",") if f.strip()]
            df = df[~df["Fecha"].apply(lambda x: datetime.strptime(x,"%d/%m/%Y").date() in fechas)]

        self.df_procesado = df.reset_index(drop=True)
        self.mostrar_en_tabla(self.df_procesado)

    # --------------------------
    # Exportar Excel con fila de totales en amarillo
    # --------------------------
    def exportar_excel(self):
        if self.df_procesado is None or self.df_procesado.empty:
            QMessageBox.warning(self, "Atención", "No hay datos para exportar.")
            return

        archivo, _ = QFileDialog.getSaveFileName(
            self, "Guardar Excel", "", "Archivos Excel (*.xlsx)"
        )
        if not archivo:
            return

        try:
            with pd.ExcelWriter(archivo, engine="openpyxl") as writer:
                # Hoja 1: Datos procesados
                df_export = self.df_procesado.copy()

                # Crear fila total por empleado dentro de la hoja
                df_final = pd.DataFrame()
                for nombre, grupo in df_export.groupby("Nombre"):
                    df_final = pd.concat([df_final, grupo], ignore_index=True)
                    # Calcular total horas extras
                    total_segundos = 0
                    for val in grupo["Horas Extra"]:
                        if val and val != "0:00:00":
                            h, m, s = map(int, val.split(":"))
                            total_segundos += h*3600 + m*60 + s
                    if total_segundos > 0:
                        df_final = pd.concat([df_final, pd.DataFrame([{
                            "Nombre": f"TOTAL HORAS EXTRAS DE {nombre}",
                            "Fecha": "",
                            "Entrada": "",
                            "Salida": "",
                            "Estado": "",
                            "Total Horas": "",
                            "Horas Extra": str(timedelta(seconds=total_segundos))
                        }])], ignore_index=True)

                df_final.to_excel(writer, index=False, sheet_name="Datos Procesados")

                # Formato amarillo para las filas de totales
                ws = writer.sheets["Datos Procesados"]
                amarillo = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
                for i, row in enumerate(df_final.itertuples(), start=2):
                    if "TOTAL HORAS EXTRAS" in str(row.Nombre):
                        for j in range(1, len(df_final.columns)+1):
                            ws.cell(row=i, column=j).fill = amarillo

                # Hoja 2: resumen total horas extras por empleado
                resumen = df_export.groupby("Nombre")["Horas Extra"].apply(
                    lambda x: str(timedelta(
                        seconds=sum(
                            int(t.split(":")[0])*3600 + int(t.split(":")[1])*60 + int(t.split(":")[2])
                            for t in x if t and t != "0:00:00"
                        )
                    ))
                ).reset_index()
                resumen.rename(columns={"Horas Extra": "Total Horas Extra"}, inplace=True)
                resumen.to_excel(writer, index=False, sheet_name="Total Horas Extra")

            QMessageBox.information(self, "Éxito", "Archivo exportado correctamente.")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo exportar el archivo:\n{e}")


# ============================================
# ============================================
# 🟢 APLICACIÓN 2 - HORAS EXTRA ANTES DE 06 AM
# ============================================
class Aplicacion2(QWidget):
    """Procesa fichajes antes de las 06:00 AM: calcula horas extra, elimina duplicados"""
    def __init__(self):
        super().__init__()
        self.df_procesado = None
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()
        self.label = QLabel("📊 Aplicación 2: Horas antes de las 06 AM")
        self.label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self.label)

        self.boton_cargar = QPushButton("Seleccionar Excel")
        self.boton_cargar.clicked.connect(self.cargar_y_procesar_excel)
        layout.addWidget(self.boton_cargar)

        self.boton_exportar = QPushButton("Exportar resultados a Excel")
        self.boton_exportar.clicked.connect(self.exportar_excel)
        layout.addWidget(self.boton_exportar)

        self.tabla = QTableWidget()
        layout.addWidget(self.tabla)

        self.setLayout(layout)

    # ----------------------------------------------------
    # PARSEO UNIVERSAL DE FECHA/HORA (soluciona caso Acosta)
    # ----------------------------------------------------
    def parsear_fecha_hora(self, valor):
        from dateutil import parser

        if pd.isna(valor):
            return None
        
        s = str(valor).strip()

        # Normalizar AM/PM en español
        s = (
            s.replace("a. m.", "AM").replace("p. m.", "PM")
             .replace("a.m.", "AM").replace("p.m.", "PM")
             .replace("am", "AM").replace("pm", "PM")
             .replace("a m", "AM").replace("p m", "PM")
        )

        # Quitar espacios dobles
        s = " ".join(s.split())

        try:
            return parser.parse(s, dayfirst=True)
        except:
            try:
                return parser.parse(s, dayfirst=False)
            except:
                return None

    # ----------------------------------------------------
    # Carga + procesado
    # ----------------------------------------------------
    def cargar_y_procesar_excel(self):
        archivo, _ = QFileDialog.getOpenFileName(
            self, "Seleccionar archivo Excel", "", "Archivos Excel (*.xlsx *.xls)"
        )
        if not archivo:
            return

        try:
            df = pd.read_excel(archivo, engine="openpyxl")

            columnas = ["Nombre", "Marc."]
            for col in columnas:
                if col not in df.columns:
                    raise Exception(f"Falta la columna requerida: {col}")

            # 🔥 PARSEO UNIVERSAL (soluciona Acosta Cristian)
            df["Marc."] = df["Marc."].apply(self.parsear_fecha_hora)

            # Eliminar las filas imposibles
            df = df.dropna(subset=["Marc."])

            df["Fecha"] = df["Marc."].dt.date
            df = df.sort_values(["Nombre", "Marc."])

            # Primer fichaje del día por persona
            df = df.drop_duplicates(subset=["Nombre", "Fecha"], keep="first")

            # Filtrar fichajes antes de 06:00 AM
            df["Hora"] = df["Marc."].dt.time
            df = df[df["Hora"] < time(6, 0, 0)].copy()

            # Cálculo horas extra
            def calcular_extra(hora):
                limite = time(6, 0, 0)
                t1 = datetime.combine(datetime.today(), hora)
                t2 = datetime.combine(datetime.today(), limite)
                segundos = int((t2 - t1).total_seconds())
                h = segundos // 3600
                m = (segundos % 3600) // 60
                s = segundos % 60
                return f"{h:02d}:{m:02d}:{s:02d}"

            df["Horas Extra"] = df["Hora"].apply(calcular_extra)

            self.df_procesado = df.reset_index(drop=True)
            self.mostrar_en_tabla()

        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo procesar el archivo:\n{e}")

    # ----------------------------------------------------
    # Mostrar datos en la tabla
    # ----------------------------------------------------
    def mostrar_en_tabla(self):
        df = self.df_procesado
        self.tabla.setRowCount(0)
        self.tabla.setColumnCount(0)
        if df.empty:
            QMessageBox.information(self, "Atención", "No hay fichajes antes de las 06:00 AM.")
            return

        self.tabla.setColumnCount(len(df.columns))
        self.tabla.setHorizontalHeaderLabels(df.columns.astype(str).tolist())

        for i, fila in df.iterrows():
            self.tabla.insertRow(i)
            for j, valor in enumerate(fila):
                self.tabla.setItem(i, j, QTableWidgetItem(str(valor)))

        self.label.setText(f"✅ Procesados ({len(df)} registros)")

    # ----------------------------------------------------
    # Exportar Excel con totales
    # ----------------------------------------------------
    def exportar_excel(self):
        if self.df_procesado is None or self.df_procesado.empty:
            QMessageBox.warning(self, "Atención", "No hay datos para exportar.")
            return

        archivo, _ = QFileDialog.getSaveFileName(
            self, "Guardar Excel", "", "Archivos Excel (*.xlsx)"
        )
        if not archivo:
            return

        try:
            with pd.ExcelWriter(archivo, engine="openpyxl") as writer:
                df_export = self.df_procesado.copy()

                df_final = pd.DataFrame()

                for nombre, grupo in df_export.groupby("Nombre"):
                    df_final = pd.concat([df_final, grupo], ignore_index=True)

                    # Total horas extra por empleado
                    total_segundos = sum(
                        int(t.split(":")[0]) * 3600 +
                        int(t.split(":")[1]) * 60 +
                        int(t.split(":")[2])
                        for t in grupo["Horas Extra"]
                    )

                    if total_segundos > 0:
                        h = total_segundos // 3600
                        m = (total_segundos % 3600) // 60
                        s = total_segundos % 60
                        total_str = f"{h:02d}:{m:02d}:{s:02d}"

                        df_final = pd.concat([
                            df_final,
                            pd.DataFrame([{
                                "Nombre": f"TOTAL DE HORAS EXTRAS DE {nombre}",
                                "Marc.": "",
                                "Fecha": "",
                                "Hora": "",
                                "Horas Extra": total_str
                            }])
                        ], ignore_index=True)

                df_final.to_excel(writer, index=False, sheet_name="Fichajes_Pre6")

                # Resaltar totales
                ws = writer.sheets["Fichajes_Pre6"]
                amarillo = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

                for i, row in enumerate(df_final.itertuples(), start=2):
                    if "TOTAL DE HORAS EXTRAS" in str(row.Nombre):
                        for j in range(1, len(df_final.columns) + 1):
                            ws.cell(row=i, column=j).fill = amarillo

                # Hoja resumen por empleado
                resumen = df_export.groupby("Nombre")["Horas Extra"].apply(
                    lambda x: sum(
                        int(t.split(":")[0])*3600 +
                        int(t.split(":")[1])*60 +
                        int(t.split(":")[2])
                        for t in x
                    )
                ).reset_index()

                resumen["Total Horas Extra"] = resumen["Horas Extra"].apply(
                    lambda s: f"{s//3600:02d}:{(s%3600)//60:02d}:{s%60:02d}"
                )
                resumen.drop(columns=["Horas Extra"], inplace=True)

                resumen.to_excel(writer, index=False, sheet_name="Total Horas Extra")

            QMessageBox.information(self, "Éxito", "Archivo exportado correctamente.")

        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo exportar el archivo:\n{e}")




# ============================================
# 🟢 APLICACIÓN 3 - CONTROL DE PRESENTISMO
# ============================================




# ============================================
# 🌟 VENTANA PRINCIPAL
# ============================================
class VentanaPrincipal(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("📊 Sistema de Control de Fichajes")
        self.resize(1200, 700)
        layout = QVBoxLayout()

        self.tabs = QTabWidget()
        self.tabs.addTab(Aplicacion1(), "Fichajes Estándar")
        self.tabs.addTab(Aplicacion2(), "Horas antes de 06 AM")
        self.tabs.addTab(AplicacionPresentismo(), "Presentismo")

        layout.addWidget(self.tabs)

        self.setLayout(layout)

# ============================================
# 🔹 EJECUCIÓN
# ============================================
if __name__ == "__main__":
    app = QApplication(sys.argv)
    ventana = VentanaPrincipal()
    ventana.show()
    sys.exit(app.exec())
